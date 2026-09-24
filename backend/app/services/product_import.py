"""Turn a seller's stock spreadsheet into draft products.

Two phases, with the rows parked in Postgres in between:

1. parse_workbook() reads the upload, validates every row and stages it. No
   product is touched. The seller sees what they are about to get.
2. commit_batch() is called repeatedly until nothing is left, creating products
   in batches. Every row carries its own status, so a commit that dies half way
   through resumes where it stopped instead of creating duplicates.

Everything lands as a draft. The sample file this was built against has no image
column and 80% of its rows are out of stock, so nothing here is fit to publish
unreviewed -- and drafts are exempt from the plan's product cap, which is what
lets a 9,608-row import through on any plan.
"""
import io
import uuid
from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from typing import Dict, List, Optional, Tuple

import openpyxl
from sqlalchemy.orm import Session

from app.models.catalog import Product, ProductCondition, ProductStatus
from app.models.product_import import (
    DuplicateAction,
    ImportRowStatus,
    ImportStatus,
    ProductImport,
    ProductImportRow,
)
from app.models.shop import Shop
from app.services.catalog import allocate_slugs

# Guard rails on an upload. The sample file is 449 KB / 9,608 rows.
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_ROWS = 20_000
# Rows per commit call. The search_vector trigger fires per row, so larger
# batches stop paying for themselves well before they risk a timeout.
DEFAULT_BATCH_SIZE = 500
# Product.price is a String(20); this keeps a garbled cell from overflowing it.
MAX_PRICE = Decimal("99999999.99")

# Headers are matched by name, case and spacing insensitive, so a reordered or
# slightly-renamed export still imports. "Description" holds the product name
# and "Brand" is a supplier grouping -- neither column is named for what it is.
HEADER_ALIASES: Dict[str, set] = {
    "sku": {"code", "sku", "barcode", "itemcode", "productcode", "item"},
    "name": {"description", "name", "productname", "product", "itemname", "details"},
    "price": {"sellingprice", "price", "unitprice", "sellprice", "retailprice"},
    "stock_qty": {"balance", "stock", "quantity", "qty", "stockqty", "onhand", "closingbalance"},
    "brand": {"brand", "category", "group", "department", "make"},
}
REQUIRED_COLUMNS = ("name", "price")


class ImportParseError(Exception):
    """The upload is unusable as a whole -- wrong file type, no rows, no headers."""


@dataclass
class ParsedRow:
    row_number: int
    raw: dict
    sku: Optional[str] = None
    name: Optional[str] = None
    price: Optional[str] = None
    stock_qty: Optional[int] = None
    brand: Optional[str] = None
    error: Optional[str] = None

    @property
    def is_valid(self) -> bool:
        return self.error is None


@dataclass
class BatchResult:
    status: ImportStatus
    processed: int = 0
    remaining: int = 0
    created: int = 0
    updated: int = 0
    skipped: int = 0
    failed: int = 0
    errors: List[str] = field(default_factory=list)


def _utcnow() -> datetime:
    return datetime.now(timezone.utc)


def _norm_header(value) -> str:
    return "".join(str(value or "").lower().split())


def _clean(value) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _parse_price(value) -> Tuple[Optional[str], Optional[str]]:
    """Money as a 2dp string, matching how Product.price is stored.

    Everything in the sample arrives as a string ("100.00"), so this goes
    through Decimal rather than float -- a price is not something to round-trip
    through binary floating point.
    """
    text = _clean(value)
    if text is None:
        return None, "No selling price"
    text = text.replace(",", "").replace("KES", "").replace("Ksh", "").strip()
    try:
        amount = Decimal(text)
    except (InvalidOperation, ValueError):
        return None, f"Price {text!r} isn't a number"
    if amount <= 0:
        return None, f"Price must be more than zero (got {text})"
    if amount > MAX_PRICE:
        return None, f"Price {text} is implausibly large"
    return str(amount.quantize(Decimal("0.01"))), None


def _parse_stock(value) -> int:
    """Stock as a whole number. Never fails the row.

    Balances arrive as "0.00000" and 54 of the sample's rows are negative, which
    is a stocktake artefact rather than something a seller should have to fix
    before importing. A missing or unreadable balance means none in stock, and a
    negative one is clamped to zero -- both are correctable during review, and
    neither is worth dropping a product over.
    """
    text = _clean(value)
    if text is None:
        return 0
    try:
        return max(0, int(Decimal(text.replace(",", ""))))
    except (InvalidOperation, ValueError):
        return 0


def parse_workbook(data: bytes, filename: str = "") -> List[ParsedRow]:
    """Read and validate an .xlsx upload. Touches no products."""
    if len(data) > MAX_UPLOAD_BYTES:
        raise ImportParseError(
            f"That file is larger than {MAX_UPLOAD_BYTES // (1024 * 1024)}MB. "
            "Split it into a few smaller uploads."
        )
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ImportParseError(
            "That doesn't look like an Excel file. Save it as .xlsx and try again."
        )

    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)

        header = next(rows, None)
        if header is None:
            raise ImportParseError("That sheet is empty.")

        columns = _map_columns(header)
        missing = [name for name in REQUIRED_COLUMNS if name not in columns]
        if missing:
            raise ImportParseError(
                "Couldn't find a "
                + " or a ".join(f"'{_example_header(name)}'" for name in missing)
                + " column. The first row of the sheet has to name the columns — "
                "download the template to see the headers."
            )

        parsed: List[ParsedRow] = []
        seen_skus: Dict[str, int] = {}

        for offset, row in enumerate(rows):
            # +2: the header is row 1, and spreadsheets count from 1
            row_number = offset + 2
            if row is None or all(cell is None or str(cell).strip() == "" for cell in row):
                continue
            if len(parsed) >= MAX_ROWS:
                raise ImportParseError(
                    f"That file has more than {MAX_ROWS:,} rows. Split it into a few uploads."
                )
            parsed.append(_parse_row(row_number, row, columns, seen_skus))

        if not parsed:
            raise ImportParseError("That sheet has headers but no product rows.")

        return parsed
    finally:
        workbook.close()


def _map_columns(header) -> Dict[str, int]:
    """Column name -> index, first match wins."""
    columns: Dict[str, int] = {}
    for index, cell in enumerate(header):
        key = _norm_header(cell)
        if not key:
            continue
        for field_name, aliases in HEADER_ALIASES.items():
            if key in aliases and field_name not in columns:
                columns[field_name] = index
    return columns


def _example_header(field_name: str) -> str:
    return {"name": "Description", "price": "SellingPrice"}.get(field_name, field_name)


def _parse_row(
    row_number: int, row, columns: Dict[str, int], seen_skus: Dict[str, int]
) -> ParsedRow:
    def cell(field_name: str):
        index = columns.get(field_name)
        if index is None or index >= len(row):
            return None
        return row[index]

    raw = {
        field_name: _clean(cell(field_name))
        for field_name in HEADER_ALIASES
        if field_name in columns
    }
    parsed = ParsedRow(row_number=row_number, raw=raw)

    # clipped to the column widths on products / product_import_rows, so an
    # oddly long cell can't abort a whole batch on insert
    sku = _clean(cell("sku"))
    parsed.sku = sku[:100] if sku else None
    brand = _clean(cell("brand"))
    parsed.brand = brand[:255] if brand else None
    parsed.stock_qty = _parse_stock(cell("stock_qty"))

    name = _clean(cell("name"))
    if name is None:
        parsed.error = "No product name"
        return parsed
    # Product.name is 255. Nothing in the sample comes close (longest is 41), and
    # keeping a product with a clipped name beats dropping it -- the seller reads
    # every one of these as a draft before it goes live.
    parsed.name = name[:255]

    parsed.price, price_error = _parse_price(cell("price"))
    if price_error:
        parsed.error = price_error
        return parsed

    if parsed.sku:
        first_seen = seen_skus.get(parsed.sku)
        if first_seen is not None:
            parsed.error = f"Same code as row {first_seen}"
            return parsed
        seen_skus[parsed.sku] = row_number

    return parsed


def stage_import(
    db: Session,
    shop: Shop,
    filename: str,
    parsed: List[ParsedRow],
    category_id: Optional[uuid.UUID] = None,
) -> ProductImport:
    """Persist the parsed rows and resolve which already exist in this shop.

    The match is recorded now because it's a fact about the shop; what to *do*
    with a match is the seller's choice at commit time. That split is what lets
    the review screen re-split "new" against "will update" as they flip the
    duplicate setting, with no round trip.
    """
    record = ProductImport(
        shop_id=shop.id,
        filename=filename[:255] or "upload.xlsx",
        status=ImportStatus.ready,
        category_id=category_id,
        total_rows=len(parsed),
    )
    db.add(record)
    db.flush()

    matches = _match_existing(db, shop.id, [row.sku for row in parsed if row.is_valid and row.sku])

    valid = invalid = matched = 0
    for row in parsed:
        matched_id = matches.get(row.sku) if (row.is_valid and row.sku) else None
        if row.is_valid:
            valid += 1
            if matched_id:
                matched += 1
        else:
            invalid += 1

        db.add(
            ProductImportRow(
                import_id=record.id,
                row_number=row.row_number,
                raw=row.raw,
                sku=row.sku,
                name=row.name,
                price=row.price,
                stock_qty=row.stock_qty,
                brand=row.brand,
                matched_product_id=matched_id,
                status=ImportRowStatus.valid if row.is_valid else ImportRowStatus.invalid,
                error=row.error,
            )
        )

    record.valid_rows = valid
    record.invalid_rows = invalid
    record.matched_rows = matched
    db.commit()
    db.refresh(record)
    return record


def _match_existing(db: Session, shop_id: uuid.UUID, skus: List[str]) -> Dict[str, uuid.UUID]:
    """sku -> existing product id, for this shop only.

    sku carries no unique constraint (it has always been free-form, and a shop
    may already hold duplicates), so where several products share a code the
    oldest one wins. Arbitrary, but stable across re-imports, which is what
    matters: the same row updates the same product every time.
    """
    if not skus:
        return {}

    found: Dict[str, uuid.UUID] = {}
    unique_skus = list({sku for sku in skus if sku})
    for start in range(0, len(unique_skus), 1000):
        chunk = unique_skus[start : start + 1000]
        rows = (
            db.query(Product.sku, Product.id)
            .filter(Product.shop_id == shop_id, Product.sku.in_(chunk))
            .order_by(Product.created_at.asc())
            .all()
        )
        for sku, product_id in rows:
            found.setdefault(sku, product_id)
    return found


def commit_batch(
    db: Session,
    import_id: uuid.UUID,
    shop: Shop,
    on_duplicate: DuplicateAction = DuplicateAction.update_stock_price,
    category_id: Optional[uuid.UUID] = None,
    limit: int = DEFAULT_BATCH_SIZE,
) -> BatchResult:
    """Create or update products for the next batch of staged rows.

    Called in a loop until `remaining` is 0. Takes the next `limit` rows still
    marked `valid` rather than an offset -- rows change status as they're
    committed, so an offset would skip work.
    """
    # Locks the import for the duration, so a seller who double-clicks Import
    # can't have two loops handing out slugs from the same starting state.
    record = (
        db.query(ProductImport)
        .filter(ProductImport.id == import_id, ProductImport.shop_id == shop.id)
        .with_for_update()
        .first()
    )
    if record is None:
        raise LookupError("Import not found")
    if record.status in (ImportStatus.completed, ImportStatus.cancelled):
        return BatchResult(status=record.status, remaining=0)

    if category_id is not None:
        record.category_id = category_id
    record.on_duplicate = on_duplicate
    if record.started_at is None:
        record.started_at = _utcnow()
    record.status = ImportStatus.processing

    rows = (
        db.query(ProductImportRow)
        .filter(
            ProductImportRow.import_id == record.id,
            ProductImportRow.status == ImportRowStatus.valid,
        )
        .order_by(ProductImportRow.row_number.asc())
        .limit(limit)
        .all()
    )

    if not rows:
        record.status = ImportStatus.completed
        record.finished_at = _utcnow()
        db.commit()
        return BatchResult(
            status=ImportStatus.completed,
            remaining=0,
            created=record.created_count,
            updated=record.updated_count,
            skipped=record.skipped_count,
            failed=record.failed_count,
        )

    result = BatchResult(status=ImportStatus.processing)

    # Matches were resolved at parse time; re-read them now because a product
    # may have been deleted in between, in which case the row creates instead.
    matched_ids = [row.matched_product_id for row in rows if row.matched_product_id]
    existing: Dict[uuid.UUID, Product] = {}
    if matched_ids:
        for product in (
            db.query(Product)
            .filter(Product.shop_id == shop.id, Product.id.in_(matched_ids))
            .all()
        ):
            existing[product.id] = product

    to_create: List[ProductImportRow] = []
    for row in rows:
        product = existing.get(row.matched_product_id) if row.matched_product_id else None
        if product is None:
            to_create.append(row)
            continue

        if on_duplicate == DuplicateAction.skip:
            row.status = ImportRowStatus.skipped
            row.error = "Already listed under this code"
            row.product_id = product.id
            result.skipped += 1
        elif on_duplicate == DuplicateAction.create_new:
            to_create.append(row)
        else:
            # Stock and price only. Never status: re-importing a stock export
            # must not knock a live product back to draft. Never name,
            # description, images or category either -- the seller may well have
            # improved those since the last import, and the spreadsheet is not
            # the better source for them.
            product.price = row.price
            product.stock_qty = row.stock_qty or 0
            row.status = ImportRowStatus.updated
            row.product_id = product.id
            result.updated += 1

    if to_create:
        slugs = allocate_slugs(db, shop.id, [row.name or "product" for row in to_create])
        for row, slug in zip(to_create, slugs):
            # id up front so the staged row can point at the product without a
            # second round trip after the flush
            product_id = uuid.uuid4()
            db.add(
                Product(
                    id=product_id,
                    shop_id=shop.id,
                    category_id=record.category_id,
                    name=row.name,
                    slug=slug,
                    price=row.price,
                    stock_qty=row.stock_qty or 0,
                    sku=row.sku,
                    condition=ProductCondition.new,
                    status=ProductStatus.draft,
                    tags=[row.brand] if row.brand else None,
                )
            )
            row.status = ImportRowStatus.created
            row.product_id = product_id
            result.created += 1

    record.created_count += result.created
    record.updated_count += result.updated
    record.skipped_count += result.skipped
    record.failed_count += result.failed

    # Flush first so the count sees this batch's new statuses. Explicit, rather
    # than leaning on the session's autoflush being off: if that ever changed,
    # a subtract-what-we-did version of this would silently report 0 remaining
    # and strand the rest of the import.
    db.flush()

    remaining = (
        db.query(ProductImportRow)
        .filter(
            ProductImportRow.import_id == record.id,
            ProductImportRow.status == ImportRowStatus.valid,
        )
        .count()
    )

    if remaining == 0:
        record.status = ImportStatus.completed
        record.finished_at = _utcnow()

    db.commit()

    result.status = record.status
    result.processed = len(rows)
    result.remaining = remaining
    # Cumulative for the whole import, not just this batch: a caller looping
    # over batches would otherwise have to sum them itself, and would
    # double-count anything it retried.
    result.created = record.created_count
    result.updated = record.updated_count
    result.skipped = record.skipped_count
    result.failed = record.failed_count
    return result
